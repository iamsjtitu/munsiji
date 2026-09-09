"""CSV / Excel / PDF statement generation (Tally-style ledger statement)."""
import csv
import io
import os
from datetime import datetime
from html import escape
from typing import Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

IST = ZoneInfo("Asia/Kolkata")
FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
GREEN = "#047857"
RED = "#B91C1C"
INK = "#18181B"
MUTED = "#71717A"
LINE = "#E4E4E7"
STRIPE = "#FAFAFA"


def _register_fonts() -> tuple[str, str]:
    """FreeSans (bundled) has the ₹ glyph and Devanagari; Helvetica (ReportLab default) prints ₹ as a black box."""
    reg, bold = os.path.join(FONT_DIR, "FreeSans.ttf"), os.path.join(FONT_DIR, "FreeSansBold.ttf")
    if os.path.exists(reg) and os.path.exists(bold):
        if "FreeSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("FreeSans", reg))
            pdfmetrics.registerFont(TTFont("FreeSans-Bold", bold))
            pdfmetrics.registerFontFamily("FreeSans", normal="FreeSans", bold="FreeSans-Bold", italic="FreeSans", boldItalic="FreeSans-Bold")
        return "FreeSans", "FreeSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def inr(v: float, decimals: int = 2) -> str:
    """Indian digit grouping: 2807120 -> 28,07,120.00"""
    v = round(abs(float(v or 0)), decimals)
    whole = int(v)
    frac = f"{v - whole:.{decimals}f}"[1:] if decimals else ""
    s = str(whole)
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts + [tail])
    return s + frac


def _d(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(IST).strftime("%d-%m-%Y")
    except Exception:  # noqa: BLE001
        return iso[:10]


def _safe_text(s: str) -> str:
    """Neutralise spreadsheet formula injection in user-provided text."""
    s = str(s or "")
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def _is_account(meta: Optional[dict]) -> bool:
    return (meta or {}).get("kind") in ("cash", "bank")


def _bal_label(v: float, acct: bool) -> str:
    """Tally-style suffix: party → Dr (lena hai) / Cr (dena hai); cash/bank → plain (in hand), negative marked."""
    if acct:
        return "" if v >= -0.004 else " (minus)"
    if v > 0.004:
        return " Dr"
    if v < -0.004:
        return " Cr"
    return ""


def _particulars(r: dict, acct: bool) -> str:
    default = ("Jama" if r["direction"] == "debit" else "Nikla") if acct else ("Diya" if r["direction"] == "debit" else "Mila")
    part = r.get("note") or default
    via = r.get("via")
    if via and not part.startswith(via):
        part += f"  [via {via}]"
    if r.get("tags"):
        part += "  " + " ".join(f"#{t}" for t in r["tags"])
    return part


def _closing_text(stmt: dict, acct: bool, kind: str) -> str:
    bal = stmt["closing_balance"]
    if acct:
        what = "cash in hand" if kind == "cash" else "bank balance"
        return f"₹{inr(bal)} {what}" + (" (minus)" if bal < -0.004 else "")
    return f"₹{inr(bal)} " + ("lena hai" if bal > 0.004 else "dena hai" if bal < -0.004 else "(settled)")


# ------------------------------------------------------------------ CSV
def to_csv(ledger_name: str, stmt: dict, subtitle: str = "", meta: Optional[dict] = None) -> bytes:
    acct = _is_account(meta)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"Ledger: {_safe_text(ledger_name)}"])
    if subtitle:
        w.writerow([subtitle])
    w.writerow(["Date", "Particulars", "In" if acct else "Debit", "Out" if acct else "Credit", "Balance"])
    w.writerow(["", "Opening Balance", "", "", f"{stmt['opening_balance']:.2f}"])
    for r in stmt["rows"]:
        w.writerow([_d(r["entry_date"]), _safe_text(_particulars(r, acct)), f"{r['amount']:.2f}" if r["direction"] == "debit" else "",
                    f"{r['amount']:.2f}" if r["direction"] == "credit" else "", f"{r['running_balance']:.2f}"])
    w.writerow(["", "Total", f"{stmt['total_debit']:.2f}", f"{stmt['total_credit']:.2f}", f"{stmt['closing_balance']:.2f}"])
    return buf.getvalue().encode("utf-8-sig")


# ------------------------------------------------------------------ Excel
def to_xlsx(ledger_name: str, stmt: dict, subtitle: str = "", meta: Optional[dict] = None) -> bytes:
    acct = _is_account(meta)
    kind = (meta or {}).get("kind", "party")
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    thin = Side(style="thin", color="D4D4D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    num_fmt = '#,##0.00;[Red]-#,##0.00'

    ws.append([f"{ledger_name} — Ledger Statement"])
    ws["A1"].font = Font(bold=True, size=14, color=GREEN.lstrip("#"))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    info = []
    if (meta or {}).get("group_name"):
        info.append(f"Group: {meta['group_name']}")
    info.append(subtitle or "Period: all entries")
    info.append(f"Closing balance: {_closing_text(stmt, acct, kind)}")
    info.append(f"Generated by Munsiji • {datetime.now(IST).strftime('%d-%m-%Y %H:%M')}")
    for line in info:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).font = Font(color=MUTED.lstrip("#"), size=10)
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
    ws.append([])

    head = ["Date", "Particulars", "In (₹)" if acct else "Debit (₹)", "Out (₹)" if acct else "Credit (₹)", "Balance (₹)"]
    ws.append(head)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN.lstrip("#"))
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    def money_cell(row_idx: int, col: int, value: Optional[float], bold: bool = False, color: Optional[str] = None):
        c = ws.cell(row=row_idx, column=col)
        if value is None:
            c.value = None
        else:
            c.value = round(float(value), 2)
            c.number_format = num_fmt
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.font = Font(bold=bold, color=(color or INK).lstrip("#"))
        c.border = border

    ws.append(["", "Opening Balance", None, None, None])
    r0 = ws.max_row
    ws.cell(row=r0, column=2).font = Font(italic=True, color=MUTED.lstrip("#"))
    money_cell(r0, 3, None)
    money_cell(r0, 4, None)
    money_cell(r0, 5, stmt["opening_balance"], bold=True)
    for col in (1, 2):
        ws.cell(row=r0, column=col).border = border

    for r in stmt["rows"]:
        ws.append([_d(r["entry_date"]), _safe_text(_particulars(r, acct)), None, None, None])
        ri = ws.max_row
        ws.cell(row=ri, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=ri, column=2).alignment = Alignment(wrap_text=True, vertical="center")
        for col in (1, 2):
            ws.cell(row=ri, column=col).border = border
        debit = r["amount"] if r["direction"] == "debit" else None
        credit = r["amount"] if r["direction"] == "credit" else None
        money_cell(ri, 3, debit, color=(GREEN if acct else RED))
        money_cell(ri, 4, credit, color=(RED if acct else GREEN))
        money_cell(ri, 5, r["running_balance"])

    ws.append(["", "Total", None, None, None])
    rt = ws.max_row
    ws.cell(row=rt, column=2).font = Font(bold=True)
    for col in (1, 2):
        ws.cell(row=rt, column=col).border = border
        ws.cell(row=rt, column=col).fill = PatternFill("solid", fgColor="F4F4F5")
    money_cell(rt, 3, stmt["total_debit"], bold=True)
    money_cell(rt, 4, stmt["total_credit"], bold=True)
    money_cell(rt, 5, stmt["closing_balance"], bold=True)
    for col in (3, 4, 5):
        ws.cell(row=rt, column=col).fill = PatternFill("solid", fgColor="F4F4F5")
    ws.append(["", f"Closing: {_closing_text(stmt, acct, kind)}"])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True, color=GREEN.lstrip("#"))

    for i, w in enumerate([13, 46, 16, 16, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def to_pdf(ledger_name: str, stmt: dict, subtitle: str = "", meta: Optional[dict] = None) -> bytes:
    acct = _is_account(meta)
    kind = (meta or {}).get("kind", "party")
    font, font_b = _register_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
                            title=f"{ledger_name} — Ledger Statement", author="Munsiji")
    s_title = ParagraphStyle("t", fontName=font_b, fontSize=17, leading=21, textColor=colors.HexColor(INK))
    s_sub = ParagraphStyle("s", fontName=font, fontSize=9.5, leading=13, textColor=colors.HexColor(MUTED))
    s_bal = ParagraphStyle("b", fontName=font_b, fontSize=11, leading=14, textColor=colors.HexColor(GREEN if stmt["closing_balance"] >= 0 or acct else RED))
    s_cell = ParagraphStyle("c", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor(INK))
    s_cell_muted = ParagraphStyle("cm", parent=s_cell, textColor=colors.HexColor(MUTED))
    s_num = ParagraphStyle("n", parent=s_cell, alignment=TA_RIGHT)
    s_num_b = ParagraphStyle("nb", parent=s_num, fontName=font_b)
    s_head = ParagraphStyle("h", fontName=font_b, fontSize=8.5, leading=11, textColor=colors.white)
    s_head_r = ParagraphStyle("hr", parent=s_head, alignment=TA_RIGHT)

    story = [Paragraph(f"{escape(ledger_name)}", s_title), Paragraph("Ledger Statement" + (f" · {escape(meta['group_name'])}" if (meta or {}).get("group_name") else ""), s_sub)]
    story.append(Paragraph(escape(subtitle) if subtitle else "Period: all entries", s_sub))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph(f"Closing balance: {escape(_closing_text(stmt, acct, kind))}", s_bal))
    story.append(Spacer(1, 5 * mm))

    def num(v: Optional[float], style=s_num, color: Optional[str] = None, suffix: str = "") -> Paragraph:
        if v is None:
            return Paragraph("", style)
        txt = inr(v) + suffix
        return Paragraph(f'<font color="{color}">{txt}</font>' if color else txt, style)

    head = [Paragraph("Date", s_head), Paragraph("Particulars", s_head), Paragraph("In (₹)" if acct else "Debit (₹)", s_head_r),
            Paragraph("Out (₹)" if acct else "Credit (₹)", s_head_r), Paragraph("Balance (₹)", s_head_r)]
    data = [head, [Paragraph("", s_cell), Paragraph("Opening Balance", s_cell_muted), Paragraph("", s_num), Paragraph("", s_num),
                   num(stmt["opening_balance"], s_num_b, suffix=_bal_label(stmt["opening_balance"], acct))]]
    for r in stmt["rows"]:
        debit = r["amount"] if r["direction"] == "debit" else None
        credit = r["amount"] if r["direction"] == "credit" else None
        data.append([
            Paragraph(_d(r["entry_date"]), s_cell),
            Paragraph(escape(_particulars(r, acct)), s_cell),
            num(debit, color=GREEN if acct else RED),
            num(credit, color=RED if acct else GREEN),
            num(r["running_balance"], suffix=_bal_label(r["running_balance"], acct)),
        ])
    data.append([Paragraph("", s_cell), Paragraph("Total", s_num_b.clone("tl", alignment=0)), num(stmt["total_debit"], s_num_b), num(stmt["total_credit"], s_num_b),
                 num(stmt["closing_balance"], s_num_b, suffix=_bal_label(stmt["closing_balance"], acct))])

    # 182mm usable width (A4 210 − 2×14 margins)
    table = Table(data, colWidths=[22 * mm, 74 * mm, 27 * mm, 27 * mm, 32 * mm], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(GREEN)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor(LINE)),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#D4D4D8")),
        ("LINEBEFORE", (2, 0), (-1, -1), 0.3, colors.HexColor(LINE)),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F4F4F5")),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#A1A1AA")),
    ]
    for i in range(2, len(data) - 1):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(STRIPE)))
    table.setStyle(TableStyle(style))
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"{len(stmt['rows'])} entries · Dr = diya / lena, Cr = mila / dena" if not acct else f"{len(stmt['rows'])} entries · In = jama, Out = nikla", s_sub))

    gen = datetime.now(IST).strftime("%d-%m-%Y %H:%M")

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(14 * mm, 9 * mm, f"Generated by Munsiji · {gen}")
        canvas.drawRightString(A4[0] - 14 * mm, 9 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


FORMATS = {
    "csv": ("text/csv", "csv", to_csv),
    "excel": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx", to_xlsx),
    "pdf": ("application/pdf", "pdf", to_pdf),
}


def build_export(fmt: str, ledger_name: str, stmt: dict, subtitle: str = "", meta: Optional[dict] = None):
    """meta: {"kind": party|cash|bank, "group_name": str} — drives In/Out vs Debit/Credit labels and the header block."""
    fmt = "excel" if fmt in ("xlsx", "xls") else fmt
    content_type, ext, fn = FORMATS[fmt]
    data = fn(ledger_name, stmt, subtitle, meta)
    safe = "".join(c if c.isalnum() else "_" for c in ledger_name).strip("_") or "ledger"
    filename = f"{safe}_statement.{ext}"
    return filename, content_type, data
